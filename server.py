import cgi
import openpyxl
from string import Template
from wsgiref.simple_server import make_server

OPTION = """<option value="$cur">$cur</option>"""

class ExchangeRate:
    def __init__(self, filepath):
        self._filepath = filepath

    def get_currencies(self):
        wb = openpyxl.load_workbook(self._filepath)
        ws = wb.worksheets[0]
        currencies = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            currencies.add(row[0])
            currencies.add(row[1])
        return sorted(currencies)

    def get_exchange_rates(self):
        wb = openpyxl.load_workbook(self._filepath)
        ws = wb.worksheets[0]
        return [[cell.value if not isinstance(cell.value, str) else cell.value.replace(",", ".")
                 for cell in row] for row in ws.iter_rows(min_row=2)]

    def obtain_rate(self, cur1, cur2, amount):
        if cur1 == cur2:
            return amount
        for c1, c2, rate in self.get_exchange_rates():
            rate = float(rate)
            if c1 == cur1 and c2 == cur2:
                return amount * rate
            elif c1 == cur2 and c2 == cur1:
                return amount / rate
        return None  # If rate not found

    def __call__(self, environ, start_response):
        path = environ.get("PATH_INFO", "").lstrip("/")
        params = {"currencies": "", "result": ""}
        status = "200 OK"
        headers = [("Content-Type", "text/html; charset=utf-8")]

        if path == "":
            currencies = "".join([Template(OPTION).substitute(cur=cur) for cur in self.get_currencies()])
            params["currencies"] = currencies
            html_file = "templates/currencies.html"

        elif path == "exchange_rate":
            form = cgi.FieldStorage(fp=environ["wsgi.input"], environ=environ)
            cur1 = form.getfirst("from", "")
            cur2 = form.getfirst("to", "")
            amount = form.getfirst("amount", "")
            if cur1 and cur2 and amount:
                try:
                    rate = self.obtain_rate(cur1, cur2, float(amount))
                    if rate is not None:
                        params["result"] = f"{amount} {cur1} = {round(rate, 2)} {cur2}"
                    else:
                        params["result"] = "Курс не знайдено."
                except Exception as e:
                    params["result"] = f"Помилка: {e}"
                html_file = "templates/exchange_rate.html"
            else:
                status = "303 SEE OTHER"
                headers.append(("Location", "/"))
                html_file = "templates/currencies.html"

        else:
            status = "404 NOT FOUND"
            html_file = "templates/error_404.html"

        start_response(status, headers)
        with open(html_file, encoding="utf-8") as f:
            page = Template(f.read()).substitute(params)
        return [bytes(page, encoding="utf-8")]


if __name__ == "__main__":
    app = ExchangeRate("data/currencies.xlsx")
    print("=== Currency Converter running at http://127.0.0.1:8000/ ===")
    make_server("", 8000, app).serve_forever()
